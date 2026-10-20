import streamlit as st
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
import io
import re
from datetime import date

st.set_page_config(page_title="Relatórios ITS", layout="wide")

# ==========================================
# CONSTANTES
# ==========================================
MONITORES_QUALIDADE = {
    "Rodrigo Ramos Novais",
    "Carlos Alessandro Silva Paiva Junior",
    "Caio Gabriel Dos Santos Lima",
    "Maria Vitoria Da Rocha",
    "Taiane Moreira De Jesus",
    "Ivone Tavares De Oliveira",
    "Yuri Fernando Dallastra Eggert",
    "Pedro Alexandre Rodrigues Carneiro",
    "Marcus Braga Dos Santos",
    "Ana Paula Silva Dos Santos",
    "Lady Lahanna Dos Santos Sutil"
}

TODOS_CICLOS = ['Ciclo 1', 'Ciclo 2', 'Ciclo 3', 'Ciclo 4']
COLUNA_DATA_CICLO = 'Data Monitoria'
TEMPLATE_DIGITAL = 'Relatório - Digital.pptx'
TEMPLATE_VOZ = 'Relatório - Voz.pptx'
MIME_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MESES = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
         'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
ORDINAIS = {1: 'primeiro', 2: 'segundo', 3: 'terceiro', 4: 'quarto'}
PADRAO_MES_CAPA = re.compile(r'^(' + '|'.join(MESES) + r') \d{4}$')

def carregar_base(file):
    if file.name.endswith('.csv'):
        try:
            return pd.read_csv(file, sep=None, engine='python', encoding='utf-8-sig')
        except UnicodeDecodeError:
            file.seek(0)
            return pd.read_csv(file, sep=None, engine='python', encoding='latin1')
    else:
        xls = pd.ExcelFile(file)
        sheet = 'BASE BRUTA' if 'BASE BRUTA' in xls.sheet_names else xls.sheet_names[0]
        return pd.read_excel(file, sheet_name=sheet)

def substituir_parcial(texto, de, para):
    # Só troca o número inteiro: "0 monitorias" não pode casar dentro de "10 monitorias"
    padrao = r'(?<![\d,.])' + re.escape(de) + r'(?!\d)'
    return re.sub(padrao, para, texto)

# ==========================================
# CABEÇALHO
# ==========================================
st.title("Relatórios Semanais de Monitorias")
st.caption("Atualize os relatórios Digital e de Voz a partir da exportação bruta de monitorias.")

# ==========================================
# 1. BASE BRUTA
# ==========================================
st.header("1. Base bruta")
uploaded_file = st.file_uploader(
    "Exportação de monitorias (.csv ou .xlsx)",
    type=["csv", "xlsx"],
    label_visibility="collapsed"
)

if uploaded_file:
    df_raw = carregar_base(uploaded_file)

    if COLUNA_DATA_CICLO not in df_raw.columns:
        st.error(f"A base não tem a coluna '{COLUNA_DATA_CICLO}'. Verifique a exportação.")
        st.stop()

    st.caption(f"{len(df_raw):,} linhas carregadas · arquivo: {uploaded_file.name}".replace(",", "."))

    # ==========================================
    # 2. CICLOS
    # ==========================================
    st.header("2. Ciclos a atualizar")
    st.caption("Marque os ciclos que serão atualizados e informe o período de cada um. Ciclos desmarcados não aparecem nos gráficos.")

    cols_ciclos = st.columns(4)
    config_ciclos = {}

    for idx, nome_ciclo in enumerate(TODOS_CICLOS):
        with cols_ciclos[idx]:
            with st.container(border=True):
                st.markdown(f"**{nome_ciclo}**")
                ativo = st.checkbox("Atualizar", value=(idx == 0), key=f"chk_{idx}")
                dt_ini = st.date_input("Início", value=date.today(), disabled=not ativo, key=f"ini_{idx}", format="DD/MM/YYYY")
                dt_fim = st.date_input("Fim", value=date.today(), disabled=not ativo, key=f"fim_{idx}", format="DD/MM/YYYY")
                config_ciclos[nome_ciclo] = {'ativo': ativo, 'ini': dt_ini, 'fim': dt_fim}

    # ==========================================
    # 3. PROCESSAMENTO
    # ==========================================
    st.header("3. Gerar relatórios")

    if st.button("Gerar relatórios", type="primary", use_container_width=True):
        ciclos_ativos = [c for c in TODOS_CICLOS if config_ciclos[c]['ativo']]
        if not ciclos_ativos:
            st.warning("Marque pelo menos um ciclo para atualizar.")
            st.stop()

        with st.spinner("Processando a base e atualizando as apresentações..."):
            df = df_raw.copy()

            # 0) Exclui monitorias lançadas na ficha "Monitoras Adicionais" (não entram no resultado)
            qtd_adicionais = 0
            if 'Questionario' in df.columns:
                eh_adicional = df['Questionario'].astype(str).str.strip().str.lower() == 'monitoras adicionais'
                qtd_adicionais = int(eh_adicional.sum())
                df = df[~eh_adicional].copy()

            # A) Identificação das colunas (por nome, com fallback para a posição original do Excel)
            col_nota = 'Nota' if 'Nota' in df.columns else df.columns[4]                 # Coluna E (índice 4)
            col_monitor = 'Monitor' if 'Monitor' in df.columns else df.columns[16]       # Coluna Q (índice 16)
            col_sla = 'Detalhe SLA' if 'Detalhe SLA' in df.columns else df.columns[-1]   # Última coluna
            col_id = 'ID' if 'ID' in df.columns else df.columns[0]                       # Coluna A (ID)

            # Coluna de Produto
            col_prod = next((c for c in df.columns if 'produto' in c.lower() or 'canal' in c.lower() or 'ilha' in c.lower()), 'Produto')

            # B) Tratamento da Nota e Regra de Classificação: =SE(E2=0;"NCG";SE(E2=100;"Excelência";"NC"))
            df['Nota_Num'] = pd.to_numeric(df[col_nota].astype(str).str.replace(',', '.'), errors='coerce').fillna(0)

            def classificar_nota(nota):
                if nota == 0:
                    return "NCG"
                elif nota == 100:
                    return "Excelência"
                else:
                    return "NC"

            df['Classificação'] = df['Nota_Num'].apply(classificar_nota)

            # C) Regra de Time (Qualidade vs Operação, pelo nome do monitor)
            monitores_norm = {m.strip().lower() for m in MONITORES_QUALIDADE}
            df['Time'] = df[col_monitor].astype(str).str.strip().str.lower().apply(
                lambda x: "Qualidade" if x in monitores_norm else "Operação"
            )

            # D) Regra de Detalhe SLA (Substituir "Vence Hoje" por "No Prazo")
            df['Detalhe SLA'] = df[col_sla].astype(str).str.strip().replace("Vence Hoje", "No Prazo")

            # E) Mapeamento de Ciclos por Intervalo de Data (sempre pela Data Monitoria)
            df['_data_lancamento'] = pd.to_datetime(df[COLUNA_DATA_CICLO], dayfirst=True, errors='coerce').dt.date
            df['Ciclo'] = None

            for nome_ciclo, cfg in config_ciclos.items():
                if cfg['ativo']:
                    mascara = (df['_data_lancamento'] >= cfg['ini']) & (df['_data_lancamento'] <= cfg['fim'])
                    df.loc[mascara, 'Ciclo'] = nome_ciclo

            # Filtra apenas as linhas que caíram nos ciclos ativos selecionados
            df = df[df['Ciclo'].notna()].copy()

            # Grupo de SLA usado nos gráficos de cumprimento (Em Abono = Justificado)
            df['SLA_Grupo'] = df['Detalhe SLA'].str.lower().map({
                'no prazo': 'No Prazo',
                'fora do prazo': 'Fora do Prazo',
                'em abono': 'Justificado'
            }).fillna('Outro')

            # F) Padronização dos Nomes de Produtos
            prod_map = {
                'CRC': 'Serasa - CRC',
                '0800 Premium': 'Serasa - 0800 Premium',
                'Chat': 'Serasa - CHAT',
                'Central de Ajuda': 'Serasa - Central de Ajuda',
                'Cadastro Positivo': 'SERASA - Cadastro Positivo',
                'Reclame Aqui': 'Serasa - Reclame Aqui',
                'GOV': 'SERASA - GOV'
            }
            df['Produto_Padrao'] = df[col_prod].astype(str).str.strip().map(prod_map).fillna(df[col_prod].astype(str).str.strip())

            # G) Tratamento de Feedback Aplicado
            col_fdb = next((c for c in df.columns if 'feedback' in c.lower()), None)
            if col_fdb:
                if pd.api.types.is_numeric_dtype(df[col_fdb]):
                    df['Fdb_Count'] = df[col_fdb].fillna(0).astype(int)
                else:
                    df['Fdb_Count'] = df[col_fdb].astype(str).str.strip().str.lower().isin(['1', 'sim', 'aplicado', 'true', 'ok']).astype(int)
            else:
                df['Fdb_Count'] = 0

            # ==========================================
            # FUNÇÕES DE AGREGAÇÃO (SOMENTE CICLOS ATIVOS)
            # ==========================================
            def get_volumes(prod, time_type):
                sub = df[(df['Produto_Padrao'] == prod) & (df['Time'] == time_type)]
                contagem = sub.groupby('Ciclo')[col_id].count()
                return [int(contagem.get(c, 0)) for c in ciclos_ativos]

            def get_medias(prod, time_type):
                sub = df[(df['Produto_Padrao'] == prod) & (df['Time'] == time_type)]
                medias = sub.groupby('Ciclo')['Nota_Num'].mean().round(2)
                return [float(medias.get(c, 0.0)) for c in ciclos_ativos]

            def get_nc_ncg(prod, class_type):
                sub = df[(df['Produto_Padrao'] == prod) & (df['Classificação'] == class_type)]
                contagem = sub.groupby('Ciclo')[col_id].count()
                return [int(contagem.get(c, 0)) for c in ciclos_ativos]

            def get_sla_pct(prod, class_type, grupo):
                # Percentual (0 a 1) de cada grupo de SLA dentro de NC ou NCG, por ciclo ativo
                sub = df[(df['Produto_Padrao'] == prod) & (df['Classificação'] == class_type)]
                valores = []
                for c in ciclos_ativos:
                    s = sub[sub['Ciclo'] == c]
                    total = len(s)
                    if total == 0:
                        valores.append(0.0)
                    else:
                        valores.append(round(float((s['SLA_Grupo'] == grupo).sum()) / total, 4))
                return valores

            def obter_dados_texto(prod_name):
                df_p = df[df['Produto_Padrao'] == prod_name]
                vol = len(df_p)
                fdb = int(df_p['Fdb_Count'].sum())
                med = f"{df_p['Nota_Num'].mean():.2f}".replace('.', ',') if vol > 0 else "0,00"
                just = len(df_p[df_p['SLA_Grupo'] == 'Justificado'])
                return str(vol), str(fdb), med, str(just)

            # ==========================================
            # MAPAS DE PRODUTOS E TEXTOS DOS TEMPLATES
            # ==========================================
            mapa_dig = {
                3: { 'prod': 'Serasa - Central de Ajuda', 'vol_antigos': ["43", "140", "[VOL]"], 'fdb_antigos': ["41", "137", "[FDB]"], 'med_antigos': ["85,47", "82,57", "[MED]"], 'just_antigos': ["4 monitorias", "5 monitorias", "[JUST] monitorias"] },
                6: { 'prod': 'Serasa - CHAT', 'vol_antigos': ["78", "316", "[VOL]"], 'fdb_antigos': ["76", "308", "[FDB]"], 'med_antigos': ["88,92", "83,87", "[MED]"], 'just_antigos': ["8 monitorias", "15 monitorias", "0 monitorias", "[JUST] monitorias"] },
                9: { 'prod': 'SERASA - GOV', 'vol_antigos': ["9", "33", "[VOL]"], 'fdb_antigos': ["9", "31", "[FDB]"], 'med_antigos': ["77,78", "88,97", "[MED]"], 'just_antigos': ["1 monitorias", "0 monitorias", "[JUST] monitorias"] },
                12:{ 'prod': 'Serasa - Reclame Aqui', 'vol_antigos': ["17", "65", "[VOL]"], 'fdb_antigos': ["17", "62", "[FDB]"], 'med_antigos': ["72,41", "76,11", "[MED]"], 'just_antigos': ["3 monitoria", "3 monitorias", "0 monitorias", "[JUST] monitorias"] }
            }

            mapa_voz = {
                3: { 'prod': 'Serasa - 0800 Premium', 'vol_antigos': ["63", "176", "[VOL]"], 'fdb_antigos': ["61", "173", "[FDB]"], 'med_antigos': ["81,73", "78,97", "[MED]"], 'just_antigos': ["1 monitoria", "10 monitorias", "0 feedbacks", "10 feedbacks", "[JUST] monitorias"] },
                6: { 'prod': 'Serasa - CRC', 'vol_antigos': ["76", "447", "[VOL]"], 'fdb_antigos': ["52", "407", "[FDB]"], 'med_antigos': ["74,80", "72,13", "[MED]"], 'just_antigos': ["2 monitorias", "43 monitorias", "10 monitorias", "[JUST] monitorias"] },
                9: { 'prod': 'SERASA - Cadastro Positivo', 'vol_antigos': ["7", "23", "[VOL]"], 'fdb_antigos': ["7", "21", "[FDB]"], 'med_antigos': ["90,57", "89,04", "90,523", "[MED]"], 'just_antigos': ["0 monitorias", "[JUST] monitorias"] }
            }

            # ==========================================
            # PERÍODO (MÊS E CICLOS) PARA OS TEXTOS DO TEMPLATE
            # ==========================================
            ativos_idx = [TODOS_CICLOS.index(c) + 1 for c in ciclos_ativos]
            inicio_ref = config_ciclos[ciclos_ativos[0]]['ini']
            mes_ref = MESES[inicio_ref.month - 1]
            ano_ref = inicio_ref.year

            meses_envolvidos = {config_ciclos[c]['ini'].month for c in ciclos_ativos}
            if len(meses_envolvidos) > 1:
                st.warning("Os ciclos marcados caem em meses diferentes. O texto usará o mês do primeiro ciclo.")

            ords = [ORDINAIS[i] for i in ativos_idx]
            lista_ords = ords[0] if len(ords) == 1 else ", ".join(ords[:-1]) + " e " + ords[-1]
            palavra_ciclo = 'ciclo' if len(ords) == 1 else 'ciclos'
            trecho_ciclos = f"do {lista_ords} {palavra_ciclo} de {mes_ref.lower()}"

            def atualizar_periodo(prs):
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if not shape.has_text_frame:
                            continue
                        for p in shape.text_frame.paragraphs:
                            for run in p.runs:
                                if PADRAO_MES_CAPA.match(run.text.strip()):
                                    run.text = f"{mes_ref} {ano_ref}"
                                elif 'do primeiro ciclo de setembro' in run.text:
                                    run.text = run.text.replace('do primeiro ciclo de setembro', trecho_ciclos)

            # ==========================================
            # ATUALIZAÇÃO DOS POWERPOINTS
            # ==========================================
            def processar_ppt(caminho_template, mapa_produtos):
                prs = Presentation(caminho_template)
                for slide_idx_vol, dados in mapa_produtos.items():
                    slide_idx_nc = slide_idx_vol + 1
                    prod_name = dados['prod']
                    slide_vol = prs.slides[slide_idx_vol]
                    slide_nc = prs.slides[slide_idx_nc]

                    # 1. Atualizar Gráficos (somente os ciclos ativos no eixo X)
                    for s in [slide_vol, slide_nc]:
                        for shape in s.shapes:
                            if shape.has_chart:
                                chart = shape.chart
                                title = chart.chart_title.text_frame.text.upper() if chart.has_title else ""

                                if "CUMPRIMENTO DO SLA" in title:
                                    classe = 'NCG' if title.strip().endswith('NCG') else 'NC'
                                    c_data = CategoryChartData(number_format='0%')
                                    c_data.categories = ciclos_ativos
                                    for grupo in ['No Prazo', 'Fora do Prazo', 'Justificado']:
                                        c_data.add_series(grupo, get_sla_pct(prod_name, classe, grupo))
                                    chart.replace_data(c_data)

                                elif "QUANTIDADE" in title:
                                    c_data = CategoryChartData()
                                    c_data.categories = ciclos_ativos
                                    c_data.add_series('Operação', get_volumes(prod_name, 'Operação'))
                                    c_data.add_series('Qualidade', get_volumes(prod_name, 'Qualidade'))
                                    chart.replace_data(c_data)

                                elif "MÉDIA" in title:
                                    c_data = CategoryChartData()
                                    c_data.categories = ciclos_ativos
                                    c_data.add_series('Operação', get_medias(prod_name, 'Operação'))
                                    c_data.add_series('Qualidade', get_medias(prod_name, 'Qualidade'))
                                    chart.replace_data(c_data)

                                elif "VOLUME" in title:
                                    c_data = CategoryChartData()
                                    c_data.categories = ciclos_ativos
                                    c_data.add_series('NC', get_nc_ncg(prod_name, 'NC'))
                                    c_data.add_series('NCG', get_nc_ncg(prod_name, 'NCG'))
                                    chart.replace_data(c_data)

                    # 2. Atualizar Textos (preservando formatação)
                    novo_vol, novo_fdb, nova_med, novas_just = obter_dados_texto(prod_name)

                    sub_exatas = []
                    for v_antigo in dados['vol_antigos']: sub_exatas.append((v_antigo, novo_vol))
                    for v_antigo in dados['fdb_antigos']: sub_exatas.append((v_antigo, novo_fdb))
                    for v_antigo in dados['med_antigos']: sub_exatas.append((v_antigo, nova_med))

                    sub_parciais = []
                    for v_antigo in dados['just_antigos']:
                        sub_parciais.append((v_antigo, f"{novas_just} monitorias"))

                    for s in [slide_vol, slide_nc]:
                        for shape in s.shapes:
                            if shape.has_text_frame:
                                for paragraph in shape.text_frame.paragraphs:
                                    for run in paragraph.runs:
                                        texto_limpo = run.text.strip()
                                        for de, para in sub_exatas:
                                            if texto_limpo == de:
                                                run.text = run.text.replace(de, para)
                                        for de, para in sub_parciais:
                                            run.text = substituir_parcial(run.text, de, para)

                # 3. Atualizar mês da capa/cabeçalhos e ciclos do texto de abertura
                atualizar_periodo(prs)

                output = io.BytesIO()
                prs.save(output)
                output.seek(0)
                return output

            try:
                ppt_dig_final = processar_ppt(TEMPLATE_DIGITAL, mapa_dig)
                ppt_voz_final = processar_ppt(TEMPLATE_VOZ, mapa_voz)

                # Base tratada para conferência
                buffer_excel = io.BytesIO()
                df.to_excel(buffer_excel, index=False, sheet_name='BASE BRUTA')
                buffer_excel.seek(0)

                # ==========================================
                # 4. RESULTADOS
                # ==========================================
                st.header("4. Resultados")
                m1, m2, m3 = st.columns(3)
                m1.metric("Monitorias processadas", f"{len(df):,}".replace(",", "."))
                m2.metric("Ficha 'Monitoras Adicionais' excluída", f"{qtd_adicionais:,}".replace(",", "."))
                m3.metric("Ciclos atualizados", len(ciclos_ativos))
                st.caption(f"Período: {mes_ref} {ano_ref} · Ciclos: " + ", ".join(ciclos_ativos))

                st.subheader("Downloads")
                d1, d2, d3 = st.columns(3)
                with d1:
                    st.download_button(
                        "Relatório Digital (.pptx)",
                        data=ppt_dig_final,
                        file_name="Relatório_Digital_Atualizado.pptx",
                        mime=MIME_PPTX,
                        use_container_width=True
                    )
                with d2:
                    st.download_button(
                        "Relatório de Voz (.pptx)",
                        data=ppt_voz_final,
                        file_name="Relatório_Voz_Atualizado.pptx",
                        mime=MIME_PPTX,
                        use_container_width=True
                    )
                with d3:
                    st.download_button(
                        "Base tratada (.xlsx)",
                        data=buffer_excel,
                        file_name="Base_Monitorias_Tratada.xlsx",
                        mime=MIME_XLSX,
                        use_container_width=True
                    )

            except Exception as e:
                st.error(f"Erro ao processar os templates: {e}")
